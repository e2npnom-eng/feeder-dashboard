#!/usr/bin/env python3
"""รวม data/**/*.xlsx (เหตุการณ์ T/R, T/L) -> data.js   ใช้: python convert.py [โฟลเดอร์] [data.js]"""
import glob, json, os, re, sys, datetime as dt
from openpyxl import load_workbook

SRC = sys.argv[1] if len(sys.argv) > 1 else "data"
OUT = sys.argv[2] if len(sys.argv) > 2 else "data.js"

A = {  # คอลัมน์ -> ชื่อหัวที่รองรับ
    "t": ["dateoutage", "date", "datetime", "วันที่", "วันเวลา"],
    "f": ["equiptment", "equipment", "feeder", "อุปกรณ์", "ฟีดเดอร์"],
    "k": ["tripreclose", "type", "ประเภท"],
    "r": ["reason", "สาเหตุ", "remark"],
}
KEY = {a: k for k, v in A.items() for a in v}
TZ = dt.timezone(dt.timedelta(hours=7))
NOW = dt.datetime.now(TZ).replace(tzinfo=None)
MIN_D, MAX_D = dt.datetime(2000, 1, 1), NOW + dt.timedelta(days=2)  # กันวันที่พิมพ์ผิด (เช่น ปี 2596 → 2053)


def when(v):
    """'dd-mm-yyyy hh:mm' (พ.ศ. หรือ ค.ศ.) หรือ datetime -> datetime ค.ศ."""
    if isinstance(v, dt.datetime):
        d = v
    else:
        m = re.match(r"\s*(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})(?:\D+(\d{1,2}):(\d{2}))?", str(v or ""))
        if not m:
            return None
        dd, mm, yy, h, mi = m.groups()
        try:
            d = dt.datetime(int(yy), int(mm), int(dd), int(h or 0), int(mi or 0))
        except ValueError:
            return None
    return d.replace(year=d.year - 543) if d.year > 2400 else d


def kind(k, r):
    for x in (str(k or ""), str(r or "")[:4]):
        if re.search(r"t\s*/\s*l|lock", x, re.I):
            return "T/L"
        if re.search(r"t\s*/\s*r|reclos", x, re.I):
            return "T/R"


rows, seen, skipped, baddate, failed, empty = [], set(), 0, 0, [], []
# ข้ามไฟล์ชั่วคราวของ Excel (~$...) ที่เกิดตอนเปิดไฟล์ค้างไว้
files = sorted(f for f in glob.glob(os.path.join(SRC, "**", "*.xlsx"), recursive=True)
               if not os.path.basename(f).startswith("~$"))
for f in files:
    name = os.path.relpath(f, SRC)
    fb = re.search(r"Event_Data_(.+?)\.xlsx$", os.path.basename(f), re.I)
    got = []
    try:
        wb = load_workbook(f, read_only=True, data_only=True)
        for ws in wb.worksheets:
            it = ws.iter_rows(values_only=True)
            cols = [KEY.get(str(h or "").strip().lower()) for h in next(it, [])]
            if "t" not in cols:
                print(f"::warning::ข้ามชีต {name}/{ws.title}: ไม่พบคอลัมน์วันที่")
                continue
            for r in it:
                o = {k: c for k, c in zip(cols, r) if k and c is not None}
                if not o:
                    continue
                d, k = when(o.get("t")), kind(o.get("k"), o.get("r"))
                fd = re.sub(r"\s+", "", str(o.get("f") or (fb.group(1) if fb else ""))).upper()  # รหัสฟีดเดอร์: ตัวพิมพ์ใหญ่ ไม่มีช่องว่าง
                if not (d and k and fd):
                    skipped += 1
                    continue
                if not (MIN_D <= d <= MAX_D):
                    baddate += 1
                    print(f"::warning::{name}: วันที่ผิดปกติ {d:%Y-%m-%d} (ฟีดเดอร์ {fd}) — ข้ามแถวนี้")
                    continue
                got.append([d.strftime("%Y-%m-%dT%H:%M"), fd, k, re.sub(r"\s+", " ", str(o.get("r") or "")).strip()])
    except Exception as e:
        failed.append(name)
        print(f"::error::อ่านไฟล์ไม่ได้ {name}: {e}")
        continue
    if not got:
        empty.append(name)
    for x in got:
        key = (x[1], x[0], x[2])
        if key not in seen:
            seen.add(key)
            rows.append(x)

# กันข้อมูลหายเงียบๆ: ถ้าไฟล์ไหนอ่านไม่ได้ หรือไม่พบเหตุการณ์เลย จะไม่เขียนทับ data.js เดิม
if failed or not rows:
    print(f"::error::ไม่เขียนทับ {OUT} — ไฟล์อ่านไม่ได้ {len(failed)} ไฟล์, เหตุการณ์ที่อ่านได้ {len(rows)} รายการ")
    sys.exit(1)
if empty:
    print(f"::warning::ไฟล์ที่ไม่พบเหตุการณ์เลย {len(empty)} ไฟล์: {', '.join(empty[:10])}")

rows.sort(key=lambda x: (x[0], x[1]), reverse=True)
upd = NOW.strftime("%Y-%m-%dT%H:%M")
body = ",\n".join(json.dumps(r, ensure_ascii=False) for r in rows)
open(OUT, "w", encoding="utf-8").write(f'window.DATA={{"updated":"{upd}","rows":[\n{body}\n]}};\n')
print(f"{OUT}: {len(rows)} เหตุการณ์ จาก {len(files)} ไฟล์ (ข้ามแถวที่อ่านไม่ได้ {skipped}, วันที่ผิดปกติ {baddate})")
